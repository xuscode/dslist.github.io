#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
价格清单 Excel → JSON 转换工具
从 "My Price List - <date>.xlsx" 读取价格数据,生成 data/price_list.json

数据来源与结构:
  - 目标 JSON 是一个扁平数组,每个元素是一条产品记录。
  - 数据来自 Excel 的两个工作表,按顺序拼接:
      1. "Industry" 表 (行业价格清单)
      2. "Portfolio" 表 (组合价格清单)
  - 两张表的表头都在第 7 行,数据从第 8 行开始,末尾有页脚注释行需要排除。

字段映射 (JSON 键 -> Excel 列名):
  Portfolio Name   -> "Portfolio Name"
  Product Number   -> "Product Number"
  Product Type     -> "Product Type"
  Short Ref        -> "Short Ref"
  Product Name     -> "Product Name"
  Release          -> "Release *"
  Licensing Scheme -> "Licensing Scheme"
  价格字段 (PLC/ALC/QLC/YLC/SLC/TBL2/TBL3/ULC/XLC/ASC/QSC/YSC/TSC3/PSC)
      -> 对应同名列。
  "Industry" 表缺少 SLC / ULC / XLC 三列,这三列在 Industry 记录中一律记为 0。

使用方法:
    cd scripts
    python xlsx_to_json.py
"""

import json
import sys
from pathlib import Path

import openpyxl

# Windows 控制台默认 GBK 编码无法输出 ✓ 等字符,统一为 UTF-8
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# JSON 输出中所有字段,按顺序排列(与旧 price_list.json 的键顺序一致)
JSON_KEYS = [
    "Portfolio Name",
    "Product Number",
    "Product Type",
    "Short Ref",
    "Product Name",
    "Release",
    "Licensing Scheme",
    "PLC",
    "ALC",
    "QLC",
    "YLC",
    "SLC",
    "TBL2",
    "TBL3",
    "ULC",
    "XLC",
    "ASC",
    "QSC",
    "YSC",
    "TSC3",
    "PSC",
]

# 非价格字段的列名映射 (JSON 键 -> Excel 表头名)
TEXT_FIELD_MAP = {
    "Portfolio Name": "Portfolio Name",
    "Product Number": "Product Number",
    "Product Type": "Product Type",
    "Short Ref": "Short Ref",
    "Product Name": "Product Name",
    "Release": "Release *",
    "Licensing Scheme": "Licensing Scheme",
}

# 价格字段 (顺序与 JSON_KEYS 中一致)
PRICE_FIELDS = [
    "PLC", "ALC", "QLC", "YLC", "SLC", "TBL2", "TBL3",
    "ULC", "XLC", "ASC", "QSC", "YSC", "TSC3", "PSC",
]

# 需要排除的页脚注释行 (出现在表头第 1 列)
FOOTER_MARKERS = (
    "For ALC and YLC prices",
    "* Release indication is information only",
    "Page ",
)

HEADER_ROW = 7  # 表头所在行 (1-indexed)
DATA_START_ROW = 8


def to_number(value):
    """把单元格值转换为数值:空/零 -> 0(int),否则 -> float。"""
    if value is None:
        return 0
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return 0
        value = float(value)
    if value == 0:
        return 0
    return float(value)


def get_header_map(ws):
    """读取表头行,返回 {列名: 列索引(0-based)}。"""
    header_values = list(next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)))
    header = {}
    for idx, name in enumerate(header_values):
        if name is not None:
            header[str(name).strip()] = idx
    return header


def is_footer(row):
    """判断是否为页脚/注释行。"""
    first = row[0]
    if first is None:
        return True
    text = str(first)
    return any(text.startswith(m) for m in FOOTER_MARKERS)


def extract_sheet(ws, sheet_name):
    """从单个工作表提取记录列表。

    Industry 表缺少 SLC/ULC/XLC 三列,需补 0。
    """
    header = get_header_map(ws)
    records = []

    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if is_footer(row):
            break  # 遇到页脚行即停止(页脚在数据末尾)

        # 跳过无产品编号的空行
        portfolio = row[header.get("Portfolio Name", 0)] if "Portfolio Name" in header else None
        product_number = row[header.get("Product Number", 1)] if "Product Number" in header else None
        if portfolio is None or product_number is None:
            continue

        record = {}

        # 文本字段 (空单元格保持为 null,与旧 JSON 约定一致)
        for json_key, excel_name in TEXT_FIELD_MAP.items():
            col = header.get(excel_name)
            val = row[col] if col is not None else None
            record[json_key] = val if val is not None else None

        # 价格字段
        for price_field in PRICE_FIELDS:
            col = header.get(price_field)
            if col is None:
                # Industry 表缺少 SLC/ULC/XLC
                record[price_field] = 0
            else:
                record[price_field] = to_number(row[col])

        records.append(record)

    return records


def main():
    script_dir = Path(__file__).parent
    project_root = script_dir.parent
    data_dir = project_root / "data"

    # 查找最新的价格清单 Excel 文件
    candidates = sorted(
        data_dir.glob("My Price List*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        print("错误: 未在 data/ 目录下找到 My Price List*.xlsx 文件")
        return False

    xlsx_path = candidates[0]
    output_file = data_dir / "price_list.json"

    print("=" * 70)
    print("价格清单 Excel → JSON 转换工具")
    print("=" * 70)
    print(f"输入文件: {xlsx_path.name}")
    print(f"输出文件: {output_file.name}")
    print("=" * 70)

    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)

    if "Industry" not in wb.sheetnames or "Portfolio" not in wb.sheetnames:
        print("错误: Excel 中缺少 'Industry' 或 'Portfolio' 工作表")
        return False

    industry_records = extract_sheet(wb["Industry"], "Industry")
    portfolio_records = extract_sheet(wb["Portfolio"], "Portfolio")

    all_records = industry_records + portfolio_records

    print(f"Industry 表记录数: {len(industry_records)}")
    print(f"Portfolio 表记录数: {len(portfolio_records)}")
    print(f"合计记录数: {len(all_records)}")

    # 写入 JSON
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(all_records, f, ensure_ascii=False, indent=2)

    print(f"\n✓ 成功生成: {output_file}")
    print(f"  共 {len(all_records)} 条记录")
    return True


if __name__ == "__main__":
    main()
